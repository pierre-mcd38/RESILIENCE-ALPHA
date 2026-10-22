#!/usr/bin/env python3
"""
PROJECT RESILIENCE-ALPHA — Diaporama complet (dossier de présentation détaillé, 16:9).

Toutes les valeurs chiffrées proviennent des fichiers de preuve du dépôt :
- ml/artifacts/model_metadata.json            (métriques ML, jeux de données, comparaison)
- data/rehearsal_runs.json                    (3 répétitions chronométrées locales)
- data/live_2pc_validation.json               (tir réel depuis le 2e PC, 19/10/2026)
- data/live_ollama_validation.json            (validation RAG / Ollama)
- data/live_storage_persistence_validation.json
- data/test_summary.json                      (21 tests unitaires/intégration)
- docs/governance/*.md, docs/known-limitations.md

Ce support ne remplace pas le support oral de 5 diapositives (Workshop2026-M2-G2-Pres.pptx).
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT_DIR / "deliverables" / "Workshop2026-M2-G2"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "Workshop2026-M2-G2-Diaporama-Complet.pptx"

BG = RGBColor(11, 19, 41)
CARD = RGBColor(22, 33, 62)
BORDER = RGBColor(46, 64, 102)
WHITE = RGBColor(241, 245, 249)
MUTED = RGBColor(148, 163, 184)
CYAN = RGBColor(56, 189, 248)
GREEN = RGBColor(52, 211, 153)
RED = RGBColor(248, 113, 113)
AMBER = RGBColor(251, 191, 36)
PURPLE = RGBColor(196, 181, 253)

TOTAL_SLIDES = 22
_state = {"n": 0}


# ----------------------------------------------------------------------------- helpers
def new_slide(prs, kicker, title):
    _state["n"] += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = kicker.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.32), Inches(1.2), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()

    ft = s.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.3))
    fp = ft.text_frame.paragraphs[0]
    fp.text = f"PROJECT RESILIENCE-ALPHA — Groupe 2 — Workshop EPSI Bac+5 2026      {_state['n']} / {TOTAL_SLIDES}"
    fp.font.size = Pt(9)
    fp.font.color.rgb = MUTED
    fp.alignment = PP_ALIGN.RIGHT
    return s


def _runs(p, item, size, color=WHITE, lead_color=None):
    if isinstance(item, tuple):
        lead, rest = item
        r1 = p.add_run()
        r1.text = lead + " "
        r1.font.bold = True
        r1.font.size = Pt(size)
        r1.font.color.rgb = lead_color or WHITE
        r2 = p.add_run()
        r2.text = rest
        r2.font.size = Pt(size)
        r2.font.color.rgb = MUTED
    else:
        r = p.add_run()
        r.text = item
        r.font.size = Pt(size)
        r.font.color.rgb = color


def bullets(s, x, y, w, h, items, size=13, bullet="•", spacing=6, lead_color=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        if bullet:
            rb = p.add_run()
            rb.text = bullet + " "
            rb.font.size = Pt(size)
            rb.font.color.rgb = CYAN
        _runs(p, item, size, lead_color=lead_color)
    return tb


def card(s, x, y, w, h, title, color=CYAN, items=None, size=12, dashed=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = color
    shp.line.width = Pt(1.5)
    if dashed:
        shp.line.dash_style = 4
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.36), Inches(0.45))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(size + 2)
    p.font.bold = True
    p.font.color.rgb = color
    if items:
        bullets(s, x + 0.18, y + 0.58, w - 0.36, h - 0.7, items, size=size, spacing=4)
    return shp


def node(s, x, y, w, h, title, sub="", color=CYAN, dashed=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.12
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = color
    shp.line.width = Pt(2)
    if dashed:
        shp.line.dash_style = 4
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(10)
        r2.font.color.rgb = WHITE
    return shp


def arrow(s, x1, y1, x2, y2, color=CYAN, label=None, lx=None, ly=None, lw=2.2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    if label:
        tb = s.shapes.add_textbox(Inches(lx if lx is not None else (x1 + x2) / 2), Inches(ly if ly is not None else (y1 + y2) / 2), Inches(lw), Inches(0.35))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = label
        p.font.size = Pt(9.5)
        p.font.italic = True
        p.font.color.rgb = color
    return c


def table(s, x, y, w, rows, col_w, size=11, header_color=CYAN, row_h=0.42):
    nrows, ncols = len(rows), len(rows[0])
    gt = s.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gt.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(16, 29, 66) if i == 0 else (CARD if i % 2 else RGBColor(18, 27, 52))
            cell.margin_left = cell.margin_right = Inches(0.08)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(size)
            p.font.bold = i == 0 or j == 0
            p.font.color.rgb = header_color if i == 0 else (WHITE if j == 0 else MUTED)
    return gt


def kpi(s, x, y, w, h, label, value, sub, color):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.1
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = color
    shp.line.width = Pt(1.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (txt, sz, col, bold) in enumerate([(label, 10, MUTED, True), (value, 26, color, True), (sub, 10, WHITE, False)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = txt
        p.font.size = Pt(sz)
        p.font.bold = bold
        p.font.color.rgb = col


def banner(s, x, y, w, h, text, color=CYAN, size=13, italic=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.2
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(16, 29, 66)
    shp.line.color.rgb = color
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.italic = italic
    p.font.color.rgb = WHITE


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------------------- deck
def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 ── Titre ---------------------------------------------------------------
    s = new_slide(prs, "Workshop national EPSI Bac+5 — Octobre 2026", "PROJECT RESILIENCE-ALPHA")
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(12), Inches(1.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = "Plateforme de paiement fictive capable de détecter une attaque applicative, de se remédier sans intervention humaine et de prouver sa reprise par une transaction métier réelle."
    p.font.size = Pt(20)
    p.font.color.rgb = WHITE
    banner(s, 0.6, 3.25, 12.1, 0.8,
           "Chaîne démontrée : T0 transaction réussie → T1 attaque depuis le PC Attaquant → T2 dégradation → T3 détection ML → T4 remédiation bornée → T5 nouvelle transaction vérifiée",
           color=GREEN, size=13)
    card(s, 0.6, 4.4, 5.9, 2.4, "Groupe 2", CYAN, [
        "Lenny BARRESI", "Youssef HANSIL", "Pierre MICHAUD-EYMOZ", "Rohail ARSHAD"], size=13)
    card(s, 6.8, 4.4, 5.9, 2.4, "Domaines du workshop couverts", AMBER, [
        "Infrastructure & automatisation", "Cybersécurité offensive / défensive",
        "IA : détection ML + assistant RAG local", "Gouvernance : EBIOS RM, PCA/PRA, NIS2/DORA",
        "IoT : authentification biométrique + ESP8266"], size=12)
    notes(s, "Présentation du projet et de l'équipe. Insister sur le fait que la preuve de reprise est une transaction métier, pas un simple redémarrage.")

    # 2 ── Problème métier --------------------------------------------------------
    s = new_slide(prs, "1. Contexte", "Le problème : une API de paiement saturée sans être « en panne »")
    card(s, 0.6, 1.7, 6.0, 4.9, "Ce qui se passe lors d'une attaque L7", RED, [
        ("Cible :", "un endpoint coûteux en calcul, POST /api/v1/transactions/verify-batch."),
        ("Effet :", "les workers saturent, la latence P95 explose, les paiements légitimes échouent."),
        ("Piège :", "le processus reste vivant ; un contrôle /health peut encore répondre."),
        ("Conséquence métier :", "encaissements bloqués, partenaires impactés, risque d'incohérence lors des reprises."),
    ], size=13)
    card(s, 6.9, 1.7, 5.8, 4.9, "Notre exigence de preuve", GREEN, [
        ("Refus des faux rétablissements :", "un conteneur redémarré ou un /health à 200 ne suffisent pas."),
        ("État RECOVERED uniquement si :", "authentification JWT, clé d'idempotence unique, écriture atomique, relecture et cohérence des soldes en centimes."),
        ("Données 100 % fictives :", "aucune carte, aucun IBAN, aucun secret de production."),
    ], size=13)
    notes(s, "Distinguer disponibilité technique et disponibilité métier. Notre critère de fin d'incident est métier.")

    # 3 ── Objectif / chaîne T0→T5 -----------------------------------------------
    s = new_slide(prs, "1. Objectif", "Chaîne de valeur obligatoire T0 → T5")
    steps = [("T0", "Transaction\ninitiale réussie", GREEN), ("T1", "Attaque contrôlée\ndepuis PC Attaquant", RED),
             ("T2", "Dégradation\nobservable", AMBER), ("T3", "Détection ML\nIsolation Forest", CYAN),
             ("T4", "Remédiation\nautomatique bornée", PURPLE), ("T5", "Nouvelle transaction\nvérifiée", GREEN)]
    w, gap, x0 = 1.8, 0.26, 0.6
    for i, (t, lbl, col) in enumerate(steps):
        x = x0 + i * (w + gap)
        node(s, x, 2.0, w, 1.5, t, lbl, col)
        if i < len(steps) - 1:
            arrow(s, x + w, 2.75, x + w + gap, 2.75, MUTED)
    card(s, 0.6, 4.0, 12.1, 2.7, "Ce que le démonstrateur doit prouver", CYAN, [
        ("Zero-Touch Recovery :", "aucune action humaine entre T1 et T5 ; l'automate décide et agit seul dans un périmètre borné."),
        ("Remédiation sûre :", "seul le service executor peut agir, sur une liste fermée d'actions signées HMAC."),
        ("Preuve métier :", "la reprise est actée par une transaction persistée et un invariant comptable vérifié."),
        ("Traçabilité :", "chaque décision et action est journalisée (chronologie SOC + journal d'audit chaîné SHA-256)."),
    ], size=13)
    notes(s, "C'est la séquence que nous montrerons en live.")

    # 4 ── Banc d'essai physique --------------------------------------------------
    s = new_slide(prs, "2. Environnement", "Banc d'essai physique à deux PC sur réseau isolé")
    node(s, 0.6, 2.2, 2.8, 1.6, "PC Attaquant", "192.168.137.135\nconsole /attacker, script PowerShell", RED)
    node(s, 4.4, 2.2, 2.6, 1.6, "Point d'accès Wi-Fi", "SSID RESILIENCE-G2\npasserelle 192.168.137.1", AMBER)
    node(s, 8.0, 1.8, 4.7, 2.4, "PC Hôte HDV0-INF-P08", "Windows 11 Pro — Intel Core Ultra 9 185H — 31,4 Go RAM\nPlateforme : WAF :8080, Centre SOC :3300\nOllama local (gemma4:e4b)", CYAN)
    arrow(s, 3.4, 3.0, 4.4, 3.0, RED, "tirs HTTP", 3.35, 2.6, 1.2)
    arrow(s, 7.0, 3.0, 8.0, 3.0, AMBER)
    node(s, 8.0, 4.6, 2.2, 1.3, "Caméra PC", "capture du visage opérateur", PURPLE)
    node(s, 10.5, 4.6, 2.2, 1.3, "ESP8266 (ESP-12F)", "COM6 — LED bicolore", GREEN)
    arrow(s, 9.1, 4.6, 9.1, 4.2, PURPLE)
    arrow(s, 11.6, 4.6, 11.6, 4.2, GREEN)
    card(s, 0.6, 4.4, 7.0, 2.4, "Garde-fous du périmètre offensif", RED, [
        "Cibles autorisées par code : 192.168.137.1 ou 127.0.0.1 uniquement",
        "Refus du réseau campus 10.60.0.0/16, des autres groupes et d'Internet",
        "Plafonds de durée, de débit et de concurrence",
        "Arrêt d'urgence : scenarios/kill_switch.py / fichier sentinelle",
    ], size=12)
    notes(s, "Tir réel validé le 19/10/2026 depuis 192.168.137.135 (preuve : data/live_2pc_validation.json).")

    # 5 ── Architecture logique ---------------------------------------------------
    s = new_slide(prs, "2. Architecture", "Composants et flux : service métier + boucle de résilience")
    node(s, 0.5, 1.7, 2.1, 1.15, "Clients", "PC Attaquant / opérateur", MUTED)
    node(s, 3.05, 1.7, 2.0, 1.15, "waf-gateway", ":8080 — signatures + blocklist", CYAN)
    node(s, 5.6, 1.7, 2.0, 1.15, "api", ":8001 — paiements", CYAN)
    node(s, 8.15, 1.7, 2.0, 1.15, "Base transactionnelle", "SQLite WAL (live)", CYAN)
    node(s, 3.05, 4.7, 2.0, 1.15, "executor", ":8005 — seul acteur", PURPLE)
    node(s, 5.6, 4.7, 2.0, 1.15, "remediator", ":8004 — automate d'état", AMBER)
    node(s, 8.15, 4.7, 2.0, 1.15, "detector", ":8003 — Isolation Forest", AMBER)
    node(s, 10.75, 1.7, 2.0, 1.15, "auth", ":8002 — JWT HS256", CYAN)
    node(s, 10.75, 3.2, 2.0, 1.15, "frontend", ":3300 — Centre SOC", CYAN)
    node(s, 10.75, 4.7, 2.0, 1.15, "rag", ":8006 — lecture seule", PURPLE, dashed=True)
    arrow(s, 2.6, 2.27, 3.05, 2.27, CYAN)
    arrow(s, 5.05, 2.27, 5.6, 2.27, CYAN)
    arrow(s, 7.6, 2.27, 8.15, 2.27, CYAN)
    arrow(s, 7.3, 2.85, 8.8, 4.7, AMBER, "télémétrie (fenêtres 5 s)", 8.0, 3.4, 2.4)
    arrow(s, 8.15, 5.27, 7.6, 5.27, AMBER)
    arrow(s, 5.6, 5.27, 5.05, 5.27, AMBER)
    arrow(s, 4.05, 4.7, 4.05, 2.85, PURPLE, "block_ip_waf (HMAC)", 2.0, 3.55, 2.0)
    arrow(s, 6.6, 4.7, 6.6, 2.85, GREEN, "sonde métier", 6.65, 3.55, 1.5)
    banner(s, 0.5, 6.2, 12.25, 0.6,
           "Cyan : flux métier  ·  Ambre : détection/décision  ·  Violet : action  ·  Vert : vérification  ·  rag (pointillés) : aucun lien vers executor ni vers le WAF",
           color=MUTED, size=11)
    notes(s, "Le RAG est consultatif, le moteur de remédiation est déterministe. Ne pas les confondre.")

    # 6 ── Modes de déploiement ---------------------------------------------------
    s = new_slide(prs, "2. Déploiement", "Deux modes d'exécution — statut réel de chacun")
    card(s, 0.6, 1.7, 6.0, 4.6, "Mode live (utilisé pour la démonstration)", GREEN, [
        ("Lancement :", "python -u scripts/run_live_platform.py"),
        ("Exécution :", "services Flask sur l'hôte, un port par service (8001 à 8006, 8080, 3300)."),
        ("Stockage :", "SQLite en mode WAL, transactions BEGIN IMMEDIATE."),
        ("Statut :", "VÉRIFIÉ — tir 2 PC, répétitions T0→T5, tests automatisés."),
    ], size=13)
    card(s, 6.9, 1.7, 5.8, 4.6, "Mode Docker Compose (industrialisation)", AMBER, [
        ("Fichier :", "infra/compose/docker-compose.yml — 14 services (postgres, auth, api, waf-gateway, detector, remediator, executor, rag, frontend, prometheus, loki, promtail, grafana, ansible-runner)."),
        ("Réseaux :", "edge_net, app_net, data_net (interne), obs_net, control_net (interne)."),
        ("Statut :", "syntaxe validée (docker compose config -q : PASS). Démarrage complet de la pile non vérifié (STANDBY_READY)."),
    ], size=12)
    banner(s, 0.6, 6.45, 12.1, 0.5, "Conséquence : Prometheus / Loki / Grafana et PostgreSQL sont configurés mais ne font pas partie de la démonstration live vérifiée.", RED, 11)
    notes(s, "Transparence : ne pas présenter la pile Compose comme démontrée en fonctionnement.")

    # 7 ── Service métier ---------------------------------------------------------
    s = new_slide(prs, "3. Service métier", "Intégrité des paiements : ce qui est garanti par le code et testé")
    card(s, 0.6, 1.7, 6.0, 5.1, "Garanties transactionnelles", CYAN, [
        ("Authentification :", "jetons JWT HS256 émis par auth, contrôle des rôles (RBAC)."),
        ("Montants :", "entiers en centimes, aucun flottant."),
        ("Idempotence :", "en-tête Idempotency-Key ; un rejeu ne crée pas de doublon."),
        ("Atomicité :", "écriture BEGIN IMMEDIATE, débit et crédit dans la même transaction."),
        ("Invariant comptable :", "total de 2 750 000 centimes (27 500,00 EUR) sur 5 comptes, vérifié après chaque opération."),
    ], size=13)
    card(s, 6.9, 1.7, 5.8, 5.1, "Preuves disponibles", GREEN, [
        ("Tests métier :", "Auth JWT, RBAC, centimes, idempotence, persistance, invariant — PASS."),
        ("Persistance :", "sauvegarde/restauration SQLite WAL vérifiée ; invariant identique avant et après ; idempotence conservée après redémarrage — PASS."),
        ("Session opérateur :", "paiements T0/T5 du tableau de bord refusés sans JWT (401) ou avec jeton falsifié (403), acceptés avec jeton valide (201)."),
    ], size=13)
    notes(s, "Source : data/test_summary.json et data/live_storage_persistence_validation.json.")

    # 8 ── WAF et scénarios offensifs --------------------------------------------
    s = new_slide(prs, "4. Cybersécurité", "WAF à deux niveaux et scénarios d'attaque disponibles")
    card(s, 0.6, 1.7, 6.0, 5.1, "waf-gateway (moteur Python)", CYAN, [
        ("Niveau 1 — signatures statiques :", "SQLi (UNION SELECT, tautologies), XSS, Path Traversal, Command Injection."),
        ("Niveau 2 — blocage dynamique :", "dynamic_blocklist alimentée uniquement par executor, TTL de 300 s."),
        ("Profils de débit :", "normal / strict_mitigation."),
        ("IP protégées :", "liste PROTECTED_IPS non bannissable."),
        ("Note :", "des fichiers nginx/ModSecurity existent dans infra/waf mais ne sont pas utilisés en mode live."),
    ], size=12)
    table(s, 6.9, 1.75, 5.8, [
        ["Scénario", "Cible"],
        ["flood_l7", "POST /api/v1/transactions/verify-batch"],
        ["brute_force_auth", "POST /auth/v1/token"],
        ["sqli_xss_injection", "verify-batch + charge SQLi/XSS"],
        ["multi_vector", "alternance des 3 vecteurs"],
    ], [2.0, 3.8], size=11, row_h=0.5)
    banner(s, 6.9, 4.5, 5.8, 2.3,
           "Scénario principal de la démonstration : flood_l7 depuis le PC Attaquant. Charge envoyée : batch_size 30, iterations 40 000, audit_profile deep_verification — l'objectif est d'épuiser les workers, pas de saturer le réseau.",
           RED, 12)
    notes(s, "Le WAF seul ne détecte pas un flood de requêtes syntaxiquement valides : c'est le rôle du ML.")

    # 9 ── Détection ML : conception ---------------------------------------------
    s = new_slide(prs, "5. Intelligence artificielle", "Détection d'anomalies : Isolation Forest sur 7 variables")
    table(s, 0.6, 1.75, 6.2, [
        ["Variable (fenêtre de 5 s)", "Importance"],
        ["latency_p95_ms", "0,188"],
        ["error_rate_ratio", "0,177"],
        ["tx_success_ratio", "0,152"],
        ["req_rate_rps", "0,133"],
        ["expensive_endpoint_ratio", "0,129"],
        ["cpu_worker_saturation", "0,122"],
        ["unique_ip_entropy", "0,099"],
    ], [4.2, 2.0], size=11, row_h=0.46)
    card(s, 7.1, 1.7, 5.6, 2.55, "Modèle", AMBER, [
        "sklearn IsolationForest + StandardScaler",
        "200 arbres, contamination 0,02, graine fixe",
        "Seuil calibré sur decision_function : -0,1176",
        "Incident confirmé après 2 fenêtres anormales consécutives",
    ], size=12)
    card(s, 7.1, 4.4, 5.6, 2.4, "Jeux de données (générés, reproductibles)", CYAN, [
        "Entraînement : 75 sessions / 900 fenêtres, trafic normal uniquement (5 régimes)",
        "Validation : 40 sessions — Test : 48 sessions (dont 18 incidents)",
        "Séparation par session_id, aucune colonne « triche »",
    ], size=12)
    notes(s, "Les données sont synthétiques mais variées (pics légitimes, erreurs 4xx bénignes, audits planifiés).")

    # 10 ── Détection ML : résultats ---------------------------------------------
    s = new_slide(prs, "5. Intelligence artificielle", "Résultats mesurés sur le jeu de test (48 sessions)")
    kpi(s, 0.6, 1.75, 2.9, 1.5, "F1-SCORE (FENÊTRES)", "0,991", "précision 0,982 — rappel 1,0", GREEN)
    kpi(s, 3.65, 1.75, 2.9, 1.5, "SESSIONS D'INCIDENT", "18 / 18", "détectées", CYAN)
    kpi(s, 6.7, 1.75, 2.9, 1.5, "FAUSSES ALERTES", "0 / 30", "sessions normales", AMBER)
    kpi(s, 9.75, 1.75, 2.95, 1.5, "DÉLAI DE CONFIRMATION", "10 s", "2 fenêtres de 5 s", PURPLE)
    table(s, 0.6, 3.55, 12.1, [
        ["Modèle comparé (test)", "Précision", "Rappel", "F1", "FPR fenêtres normales", "Détection sessions"],
        ["IsolationForest (retenu)", "0,982", "1,000", "0,991", "0,72 %", "100 %"],
        ["OneClassSVM", "0,982", "1,000", "0,991", "0,72 %", "100 %"],
        ["EllipticEnvelope", "1,000", "0,982", "0,991", "0 %", "100 %"],
    ], [3.4, 1.5, 1.5, 1.3, 2.4, 2.0], size=11, row_h=0.45)
    banner(s, 0.6, 5.55, 12.1, 1.25,
           "Lecture honnête : les trois modèles sont équivalents sur ces données synthétiques. Ces scores ne garantissent pas les performances sur du trafic réel ; un réentraînement sur traces de production serait nécessaire.",
           AMBER, 12)
    notes(s, "Source : ml/artifacts/model_metadata.json. ROC AUC test 0,996. Inférence ~0,03 ms par fenêtre.")

    # 11 ── Remédiation -----------------------------------------------------------
    s = new_slide(prs, "6. Remédiation autonome", "Automate d'incident et exécuteur confiné")
    states = [("NORMAL", GREEN), ("SUSPECT", AMBER), ("INCIDENT", RED), ("REMEDIATING", PURPLE), ("RECOVERED", GREEN)]
    w, gap = 2.15, 0.33
    for i, (st, col) in enumerate(states):
        x = 0.6 + i * (w + gap)
        node(s, x, 1.75, w, 0.85, st, "", col)
        if i < len(states) - 1:
            arrow(s, x + w, 2.17, x + w + gap, 2.17, MUTED)
    node(s, 8.05, 2.95, 2.15, 0.75, "FAILED", "arrêt sûr après 3 échecs", RED)
    arrow(s, 9.1, 2.6, 9.1, 2.95, RED)
    card(s, 0.6, 3.95, 6.0, 2.85, "executor : liste fermée (ALLOWLIST)", PURPLE, [
        "block_ip_waf / unblock_ip_waf",
        "set_rate_limit_profile",
        "recycle_api_workers",
        "reset_connection_pool",
        "Tout autre ordre est rejeté",
    ], size=12)
    card(s, 6.9, 3.95, 5.8, 2.85, "Contrôles de sécurité (testés)", GREEN, [
        "Signature HMAC-SHA256 obligatoire",
        "Anti-rejeu : horodatage ±30 s et command_id unique",
        "Validation stricte des paramètres (Pydantic), refus de toute commande shell",
        "Journal d'audit chaîné SHA-256",
        "Socket Docker jamais monté dans les services applicatifs",
    ], size=12)
    notes(s, "Tests sécurité exécuteur : HMAC, anti-rejeu, allowlist, anti-shell, audit SHA-256 — PASS.")

    # 12 ── Preuve de reprise -----------------------------------------------------
    s = new_slide(prs, "6. Vérification", "La preuve de reprise : une transaction, pas un redémarrage")
    card(s, 0.6, 1.7, 6.0, 5.1, "Sonde métier automatique (RecoveryProof)", GREEN, [
        ("Opération :", "virement de 1,00 EUR depuis le compte ACC-PROBE-SRC, initié par resilience_probe."),
        ("Contrôles :", "API saine, persistance en base, doublon d'idempotence refusé, invariant comptable vérifié."),
        ("Exemple réel (19/10/2026) :", "incident INC-20261019-2CF85F, transaction tx-fc0319d9…, durée totale de la sonde 94 ms."),
        ("Puis T5 :", "transaction opérateur de 25,00 EUR depuis le tableau de bord authentifié."),
    ], size=13)
    table(s, 6.9, 1.75, 5.8, [
        ["Ce qui NE suffit PAS", "Pourquoi"],
        ["Conteneur redémarré", "ne prouve pas que les paiements passent"],
        ["/health à 200", "répond même sous saturation partielle"],
        ["IP attaquante bloquée", "restaure la disponibilité, pas l'intégrité"],
        ["Fin d'une compromission", "hors périmètre : exigerait isolation, révocation de clés, forensique"],
    ], [2.3, 3.5], size=11, row_h=0.75)
    notes(s, "Insister : restauration de disponibilité ≠ résolution d'une compromission.")

    # 13 ── RAG -------------------------------------------------------------------
    s = new_slide(prs, "5. Intelligence artificielle", "Assistant RAG local : consultatif, sans aucun pouvoir d'action")
    card(s, 0.6, 1.7, 6.0, 5.1, "Fonctionnement", PURPLE, [
        ("Corpus :", "documents Markdown de docs/ et AGENTS.md, découpés par sections."),
        ("Recherche :", "TF-IDF (scikit-learn), citations systématiques [Source: fichier#section]."),
        ("Génération :", "LangChain + Ollama, modèle gemma4:e4b (8,0 B, Q4_K_M), exécuté localement."),
        ("Garde-fou :", "refus de répondre sous le seuil de similarité ; mode OFFLINE_RETRIEVAL_ONLY si Ollama est arrêté."),
        ("Isolation :", "aucun outil système, aucun client vers executor."),
    ], size=12)
    table(s, 6.9, 1.75, 5.8, [
        ["Test (21/10/2026)", "Résultat"],
        ["Procédure de remédiation", "réponse citée — 11,5 s"],
        ["Idempotence du grand livre", "réponse « information absente » — 2,1 s"],
        ["Question hors sujet", "refus (sous le seuil) — PASS"],
        ["Ollama arrêté", "repli documentaire vérifié"],
    ], [2.9, 2.9], size=11, row_h=0.55)
    banner(s, 6.9, 4.75, 5.8, 2.05,
           "Limite observée : sur la question d'idempotence, l'assistant n'a pas retrouvé la bonne section alors que l'information existe dans le code. La recherche TF-IDF reste perfectible.",
           AMBER, 12)
    notes(s, "Le RAG aide l'opérateur à comprendre ; il ne décide ni n'exécute rien.")

    # 14 ── Biométrie / IoT ------------------------------------------------------
    s = new_slide(prs, "7. IoT & contrôle d'accès", "Authentification multifacteur : visage + ESP8266 + mot de passe")
    flow = [("Caméra PC", "capture navigateur", PURPLE), ("Détection visage", "Haar Cascade OpenCV", CYAN),
            ("Comparaison", "descripteur 128D, seuil 0,75", AMBER), ("LED ESP8266", "vert clignotant / rouge", GREEN),
            ("Mot de passe", "JWT émis", CYAN), ("Paiements", "Section T0/T5 déverrouillée", GREEN)]
    w, gap = 1.85, 0.2
    for i, (t, sub, col) in enumerate(flow):
        x = 0.6 + i * (w + gap)
        node(s, x, 1.75, w, 1.2, t, sub, col)
        if i < len(flow) - 1:
            arrow(s, x + w, 2.35, x + w + gap, 2.35, MUTED)
    card(s, 0.6, 3.3, 6.0, 3.5, "Ce qui est implémenté et vérifié", GREEN, [
        "Descripteur : grille de luminance 8×8 + gradients Sobel 8×8",
        "Rejet si visage absent ou trop proche (> 75 % du cadre)",
        "Essai manuel : même visage ≈ 0,87 ; autre visage ≈ 0,55 (rejeté)",
        "LED : vert clignotant (visage OK), vert fixe 30 s (accès), rouge fixe 30 s, rouge clignotant 10×",
        "Paiements refusés côté serveur sans JWT valide",
    ], size=12)
    card(s, 6.9, 3.3, 5.8, 3.5, "Limites à assumer devant le jury", RED, [
        "Descripteur artisanal, pas un réseau de neurones de reconnaissance faciale",
        "Aucune détection du vivant : une photo pourrait tromper le système",
        "Verrouillage 5 min désactivé pour les essais ; mot de passe de démonstration simplifié",
        "Seul le profil de Rohail ARSHAD est enrôlé avec un vrai visage ; les autres profils sont des références synthétiques",
    ], size=12)
    notes(s, "Rester factuel : c'est une preuve de concept IoT, pas une biométrie de niveau bancaire.")

    # 15 ── Démonstration ---------------------------------------------------------
    s = new_slide(prs, "8. Démonstration en direct", "Ce que le jury va observer")
    table(s, 0.6, 1.75, 12.1, [
        ["Étape", "Action", "Preuve visible"],
        ["Accès", "Visage + mot de passe opérateur", "LED ESP8266, badge de session vert, section paiements déverrouillée"],
        ["T0", "Transaction de 50,00 EUR", "ID de transaction COMMITTED, invariant 27 500,00 EUR"],
        ["T1", "Tir flood_l7 depuis 192.168.137.135", "Débit et IP source sur le radar IA"],
        ["T2", "Saturation", "Latence P95 et taux d'erreur en hausse"],
        ["T3", "Détection", "Score sous -0,1176, compteur 2/2, état INCIDENT"],
        ["T4", "Remédiation", "IP bannie, réponses 403 côté attaquant, état REMEDIATING"],
        ["T5", "Transaction de 25,00 EUR + sonde", "État RECOVERED, nouvelle transaction persistée"],
    ], [1.3, 4.2, 6.6], size=12, row_h=0.6)
    notes(s, "Environ 2 min 30 de live. Pas de captures : tout est montré en direct.")

    # 16 ── Mesures ---------------------------------------------------------------
    s = new_slide(prs, "9. Résultats", "Mesures réelles et leurs conditions")
    kpi(s, 0.6, 1.75, 2.9, 1.5, "MTTD (RÉPÉTITIONS)", "≈ 10,1 s", "2 fenêtres de 5 s", CYAN)
    kpi(s, 3.65, 1.75, 2.9, 1.5, "MTTR APRÈS DÉTECTION", "≈ 0,7 s", "remédiation + sonde", GREEN)
    kpi(s, 6.7, 1.75, 2.9, 1.5, "LATENCE P95", "≈ 875 → 15 ms", "dégradée → après reprise", AMBER)
    kpi(s, 9.75, 1.75, 2.95, 1.5, "TESTS AUTOMATISÉS", "21 / 21", "0 échec, 0 erreur", PURPLE)
    card(s, 0.6, 3.5, 6.0, 3.3, "3 répétitions chronométrées (22/10/2026)", CYAN, [
        "Exécutées sur l'hôte, attaquant simulé (IP 192.168.137.45)",
        "3 / 3 PASS : invariant OK, chaîne d'audit valide",
        "Durée totale d'incident ≈ 10,8 s par répétition",
        "12 requêtes attaquantes bloquées après remédiation",
    ], size=12)
    card(s, 6.9, 3.5, 5.8, 3.3, "Tir réel depuis le 2e PC (19/10/2026)", GREEN, [
        "IP 192.168.137.135 bannie (score ML -0,1875), TTL 300 s",
        "533 requêtes bloquées par la blocklist dynamique",
        "Sonde de reprise persistée, état RECOVERED",
        "Vidéo chronométrée obligatoire : non encore enregistrée",
    ], size=12)
    notes(s, "Les mesures du 19/10 (MTTD 2 s) et du 22/10 (≈10 s) proviennent de configurations différentes ; ne pas les fusionner.")

    # 17 ── EBIOS RM --------------------------------------------------------------
    s = new_slide(prs, "10. Gouvernance — EBIOS RM", "Événements redoutés et mesures du démonstrateur")
    table(s, 0.6, 1.75, 12.1, [
        ["Événement redouté", "Description", "Mesure implémentée"],
        ["ER-01", "Indisponibilité prolongée des paiements (saturation L7)", "Détection ML + blocage WAF + sonde métier"],
        ["ER-02", "Altération du grand livre / double débit", "Centimes entiers, idempotence, atomicité, invariant"],
        ["ER-03", "Détournement de la chaîne de remédiation", "ALLOWLIST, HMAC, anti-rejeu, IP protégées, kill switch"],
        ["ER-04", "Falsification du journal d'audit / des preuves", "Journal chaîné SHA-256 vérifié"],
        ["ER-05", "Divulgation de secrets ou manipulation du RAG", "Chiffrement SOPS-compatible (AES-GCM), scan de secrets, RAG sans outils"],
    ], [2.0, 5.0, 5.1], size=12, row_h=0.72)
    banner(s, 0.6, 6.2, 12.1, 0.6, "Sources de risque étudiées : attaquant externe (déni de service), fraudeur (double dépense), attaquant ciblant l'automatisation AIOps/LLM.", MUTED, 11)
    notes(s, "Source : docs/governance/ebios-rm.md (ateliers 1 à 5).")

    # 18 ── PCA / PRA -------------------------------------------------------------
    s = new_slide(prs, "10. Gouvernance — Continuité", "PCA / PRA : objectifs cibles vs mesures observées")
    table(s, 0.6, 1.75, 12.1, [
        ["Indicateur", "Objectif cible", "Observé sur le démonstrateur"],
        ["RTO", "≤ 15 s (incident L7) / ≤ 15 min (reconstruction sur hôte de secours)", "≈ 10,8 s pour le scénario L7 ; non automatisé pour une panne matérielle"],
        ["RPO", "0 (défaillance applicative) / ≤ 5 min (perte disque)", "0 transaction perdue lors des scénarios testés"],
    ], [1.6, 5.2, 5.3], size=12, row_h=0.8)
    card(s, 0.6, 4.35, 6.0, 2.45, "Couvert par la reprise autonome", GREEN, [
        "Saturation applicative L7", "Blocage d'une IP hostile", "Recyclage des workers API", "Vérification par sonde idempotente"], size=12)
    card(s, 6.9, 4.35, 5.8, 2.45, "Non couvert (escalade humaine)", RED, [
        "Panne matérielle ou corruption du disque du PC hôte unique",
        "Compromission de la base ou fuite de secrets",
        "DDoS volumétrique saturant le lien Wi-Fi en amont",
        "Automate en état FAILED après 3 échecs"], size=12)
    notes(s, "Un redémarrage de conteneur n'est pas un PCA.")

    # 19 ── NIS2 / DORA -----------------------------------------------------------
    s = new_slide(prs, "10. Gouvernance — Réglementation", "NIS2 / DORA : ce que notre travail permet d'aborder")
    card(s, 0.6, 1.7, 3.9, 4.6, "A. Démontré techniquement", GREEN, [
        "Détection d'incident TIC basée sur le comportement",
        "Réponse et reprise automatisées et bornées",
        "Traçabilité des actions (audit chaîné)",
        "Test de résilience par tir contrôlé réel",
    ], size=12)
    card(s, 4.7, 1.7, 3.9, 4.6, "B. Bonnes pratiques appliquées", CYAN, [
        "Moindre privilège et cloisonnement",
        "Chiffrement des secrets, scan anti-fuite",
        "Analyse de risques EBIOS RM",
        "Distinction objectifs RTO/RPO vs mesures",
    ], size=12)
    card(s, 8.8, 1.7, 3.9, 4.6, "C. Hors périmètre du prototype", RED, [
        "Gouvernance et responsabilité de la direction (NIS2 art. 20)",
        "Notification des incidents 24 h / 72 h / 1 mois (NIS2 art. 23, DORA chap. III)",
        "Gestion des prestataires TIC tiers (DORA chap. V)",
        "Redondance multi-sites",
    ], size=12)
    banner(s, 0.6, 6.45, 12.1, 0.5, "Le prototype n'est pas « conforme NIS2/DORA » : il illustre des mesures techniques attendues (NIS2 art. 21, DORA chap. II et IV).", AMBER, 11)
    notes(s, "Source : docs/governance/nis2-dora-compliance.md. DORA = Règlement (UE) 2022/2554 ; NIS2 = Directive (UE) 2022/2555.")

    # 20 ── Limites ---------------------------------------------------------------
    s = new_slide(prs, "11. Limites", "Limites connues, documentées sans complaisance")
    card(s, 0.6, 1.7, 6.0, 5.1, "Techniques", RED, [
        ("Hôte unique :", "toute la plateforme tourne sur un seul PC (SPOF matériel)."),
        ("Données ML synthétiques :", "scores non extrapolables à un trafic réel ; sensibilité au concept drift."),
        ("Attaques « low and slow » distribuées :", "peuvent rester sous le seuil ; le ban par IP devient inefficace (CGNAT, botnet)."),
        ("DDoS volumétrique L3/L4 :", "hors de portée d'un WAF local."),
    ], size=12)
    card(s, 6.9, 1.7, 5.8, 5.1, "Vérifications incomplètes", AMBER, [
        ("Docker Compose :", "pile complète non démarrée/vérifiée en conditions réelles."),
        ("Observabilité :", "Prometheus/Loki/Grafana configurés mais hors démonstration live."),
        ("Vidéo chronométrée :", "non enregistrée à ce jour."),
        ("Biométrie :", "pas de détection du vivant, descripteur simple."),
        ("RAG :", "recherche TF-IDF parfois imprécise."),
    ], size=12)
    notes(s, "Source : docs/known-limitations.md et data/test_summary.json.")

    # 21 ── Perspectives ----------------------------------------------------------
    s = new_slide(prs, "12. Perspectives", "Vers un environnement industriel")
    table(s, 0.6, 1.75, 12.1, [
        ["Priorité", "Amélioration", "Limite traitée"],
        ["1", "Déployer et valider la pile Docker Compose complète avec observabilité Prometheus/Grafana", "Vérifications incomplètes"],
        ["2", "Cluster multi-zones actif/actif + réplication PostgreSQL synchrone", "SPOF de l'hôte unique"],
        ["3", "WAF en bordure (CDN / Anycast), filtrage par jeton ou empreinte TLS plutôt que par IP", "DDoS volumétrique, CGNAT"],
        ["4", "Secrets dans Vault / HSM, rotation automatique", "Gestion des clés locale"],
        ["5", "MLOps : réentraînement sur traces réelles, validation humaine des actions destructrices", "Données synthétiques, drift"],
    ], [1.2, 7.2, 3.7], size=12, row_h=0.72)
    notes(s, "Feuille de route détaillée : docs/governance/executive-recommendations.md.")

    # 22 ── Conclusion ------------------------------------------------------------
    s = new_slide(prs, "Conclusion", "Ce que nous retenons")
    bullets(s, 0.6, 1.8, 12.1, 3.6, [
        ("Une boucle fermée fonctionnelle :", "détection ML, remédiation signée et confinée, preuve de reprise par transaction."),
        ("Une exigence de preuve métier :", "l'état RECOVERED n'est jamais déclaré sur un simple redémarrage."),
        ("Une IA à sa juste place :", "Isolation Forest détecte, l'automate décide, executor agit, le RAG ne fait que conseiller."),
        ("Une transparence assumée :", "limites et vérifications manquantes documentées, aucune conformité réglementaire revendiquée."),
    ], size=17, spacing=14)
    banner(s, 0.6, 5.6, 12.1, 1.1, "Merci — Questions ?", CYAN, 26)
    notes(s, "Ouvrir sur les questions du jury.")

    assert _state["n"] == TOTAL_SLIDES, _state["n"]
    prs.save(OUT_PATH)
    print(f"[OK] {OUT_PATH} — {TOTAL_SLIDES} diapositives — {OUT_PATH.stat().st_size} octets")


if __name__ == "__main__":
    build()
